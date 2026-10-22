# -*- coding: utf-8 -*-
"""IoT Sentinel - presentation / viva-voce guide (Word)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'IoT_Sentinel_Viva_Guide.docx')
d = Document()
for s in d.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(72)


def style_fonts():
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet'):
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.font.color.rgb = RGBColor(0, 0, 0)
    d.styles['Normal'].font.size = Pt(12)
    d.styles['Heading 1'].font.size = Pt(16)
    d.styles['Heading 2'].font.size = Pt(14)
    d.styles['Heading 3'].font.size = Pt(12)
    for n in ('Heading 1', 'Heading 2', 'Heading 3'):
        d.styles[n].font.bold = True


style_fonts()


def h1(t): d.add_heading(t, 1)
def h2(t): d.add_heading(t, 2)
def h3(t): d.add_heading(t, 3)


def p(t, bold_lead=None):
    para = d.add_paragraph()
    para.alignment = AL.JUSTIFY
    para.paragraph_format.space_after = Pt(6)
    if bold_lead:
        r = para.add_run(bold_lead + ' '); r.bold = True
    para.add_run(t)
    return para


def bl(items):
    for it in items:
        para = d.add_paragraph(style='List Bullet')
        para.alignment = AL.JUSTIFY
        if isinstance(it, tuple):
            r = para.add_run(it[0] + ' '); r.bold = True
            para.add_run(it[1])
        else:
            para.add_run(it)


def qa(q, a):
    para = d.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.keep_with_next = True
    r = para.add_run('Q. ' + q); r.bold = True
    ans = d.add_paragraph()
    ans.alignment = AL.JUSTIFY
    ans.paragraph_format.space_after = Pt(8)
    ans.add_run('A. ' + a)


t = d.add_paragraph(); t.alignment = AL.CENTER
r = t.add_run('IoT SENTINEL'); r.bold = True; r.font.size = Pt(20)
t = d.add_paragraph(); t.alignment = AL.CENTER
r = t.add_run('IoT Device Health and Security Monitoring Platform with AI Maintenance Assistant'); r.bold = True; r.font.size = Pt(14)
t = d.add_paragraph(); t.alignment = AL.CENTER
t.add_run('Presentation and viva-voce guide: what it is, why it was built, how every part works, and the questions you will be asked').italic = True

# ------------------------------------------------------------------------------------------------------------------
h1('1. The project in one minute')
p('IoT Sentinel is a web platform that watches a fleet of connected devices (pumps, CNC machines, gateways, cameras, sensors, meters, trackers), '
  'tells you when something is wrong, warns you BEFORE a limit is crossed, and lets you ask questions in plain English and get answers with sources. '
  'Everything runs on your own computer: nothing is sent to an outside AI service.')
p('Use this as your opening: "Maintenance teams usually find out about a problem only after a fixed limit has been crossed, and the knowledge needed to understand it sits in '
  'manuals nobody has time to search. IoT Sentinel monitors device health and security, predicts an imminent limit breach and explains why, and answers maintenance '
  'questions from live data and uploaded documents, always showing the source."')

h1('2. Problem statement, objectives, scope')
h2('2.1 Problem statement')
bl([
    ('Reactive monitoring.', 'Normal alert systems fire only after a threshold is crossed - the machine is already in trouble.'),
    ('One size does not fit all.', 'A pump, a CNC machine and a battery sensor behave differently, so one fixed limit for everything is wrong.'),
    ('Silent failure.', 'A device that stops sending data raises no alert unless something notices the silence (heartbeat).'),
    ('Scattered knowledge.', 'Manuals and procedures are in PDFs and Word files; finding the right paragraph during a fault is slow.'),
    ('Security and accountability.', 'Firmware versions, security incidents, maintenance history and who changed what are usually tracked in different places or not at all.'),
])
h2('2.2 Objectives')
bl([
    'Monitor eight signals per device (temperature, vibration, pressure, motor current, CPU, memory, battery, signal) with limits that can be overridden per machine type.',
    'Detect offline devices through expiring heartbeats.',
    'Predict, for every signal, the probability that its limit will be crossed within the next 20 readings, and explain the prediction exactly.',
    'Track firmware, security incidents, maintenance history, notifications and an audit trail.',
    'Provide an AI Maintenance Assistant that answers only from live data and uploaded documents, with verifiable citations, and that can never change anything.',
    'Run locally and reproducibly (Docker), with automated tests.',
])
h2('2.3 Scope and honest limits (say these yourself - examiners respect it)')
bl([
    'The demo fleet is produced by a built-in simulator (18 devices, 8 machine types), not real hardware. Real devices can send readings through the same REST API.',
    'The prediction model reaches AUC about 0.94, precision about 89 percent, recall about 73 percent on held-out SYNTHETIC episodes. These numbers describe the generator, not real machines, and the UI labels the model source as "synthetic".',
    'A risk percentage means "probability the configured limit is crossed within 20 readings". It is NOT the probability that equipment will fail.',
    'Scanned (image-only) PDFs are not read because there is no OCR.',
    'The AI assistant uses a small local model; it can be wrong, which is why every answer shows its sources and the system refuses to answer when sources do not support it.',
])

# ------------------------------------------------------------------------------------------------------------------
h1('3. Architecture and workflow')
h2('3.1 The big picture')
p('Five containers run together with Docker Compose: frontend (React app served by nginx), backend (Node.js + Express API), MongoDB (operational data), '
  'PostgreSQL with pgvector (knowledge base and chat), and Redis (fast in-memory state). The language model runs OUTSIDE Docker on the host computer, in Ollama, '
  'because it needs the graphics card.')
bl([
    ('Browser (React)', 'talks to the REST API; receives live updates and streamed AI answers.'),
    ('Backend (Express)', 'authentication, validation, ingestion, alert engine, prediction, reports, assistant orchestration, audit log.'),
    ('MongoDB', 'devices, telemetry, alerts, rules, users, firmware, incidents, maintenance, prediction history, model versions, audit.'),
    ('PostgreSQL + pgvector', 'knowledge-base documents, their text chunks and embedding vectors, chat conversations and messages.'),
    ('Redis', 'rolling window of the last 20 readings per signal, rate limiting, locks (one answer at a time per chat).'),
    ('Ollama (host)', 'runs the chat model and the embedding model locally.'),
])
h2('3.2 Workflow A: a reading travels through the system')
bl([
    '1. A device (or the simulator) sends a reading with its own API key: POST /api/telemetry. The key is stored only as a hash.',
    '2. The backend validates the data (zod schemas) and stores it in MongoDB; the latest values go into the Redis rolling window.',
    '3. The alert engine compares each signal with the limits for that machine type (default or overridden). A breach opens an alert; a recovery closes it.',
    '4. The prediction engine takes the last 20 readings, fits a line, builds five features and runs the logistic regression: probability of breaching within 20 readings.',
    '5. If the probability is above 50 percent and the linear time-to-breach agrees, a PREDICTIVE alert opens with an estimated time. Otherwise the signal is only marked "elevated".',
    '6. Webhooks and notifications are sent, the event is written to the audit/event stream, and the dashboard updates live.',
    '7. A background sweeper marks devices offline when their heartbeat expires.',
])
h2('3.3 Workflow B: a question to the AI assistant')
bl([
    '1. The user types a question in the Assistant page. The backend saves it and takes a per-chat lock.',
    '2. A deterministic rule-based planner (router.ts) decides which read-only tools are needed (fleet status, top risk, explain prediction, history, and so on). The language model does not choose tools.',
    '3. Each tool runs fixed, parameterised queries, with a timeout and bounded rows, and is written to the audit log.',
    '4. If the question needs written guidance, the knowledge base is searched: the question is embedded, then vector search and keyword search are merged.',
    '5. Only excerpts the user is allowed to read are retrieved (private files stay private); weak matches are dropped by a relevance gate.',
    '6. If there is no usable source, a fixed reply says so and the model is not even called.',
    '7. Otherwise the sources are placed in a prompt between untrusted markers and the model answers, streaming tokens to the screen.',
    '8. The answer is checked: citations such as [D1] or [S2] must refer to real sources; invented citations are removed; the UI shows "cited versus considered" sources.',
])

# ------------------------------------------------------------------------------------------------------------------
h1('4. Technology choices and why')
h2('4.1 What is React and why do we use it?')
p('React is a JavaScript library for building user interfaces out of small reusable pieces called components (a table, a chart, a sidebar). You describe what the screen should look like '
  'for a given state, and React updates only the parts of the page that changed. That matters here because a monitoring dashboard changes constantly: new readings, alerts, '
  'predictions and streamed AI words arrive every moment, and re-drawing the whole page each time would be slow and flickery.')
bl([
    'Reusable components: the same device card, risk badge and chart are used on many pages.',
    'Fast updates through the virtual DOM: only changed elements are redrawn - ideal for live data.',
    'Large ecosystem: React Router (pages), TanStack Query (fetching, caching, refreshing data), Recharts (charts).',
    'TypeScript on top gives type checking, so mistakes are found while writing, not in front of the user.',
    'Vite is the build tool: instant dev server and optimised production bundle. In Docker the built files are served by nginx.',
])
h2('4.2 Node.js and Express')
p('Node.js runs JavaScript on the server, so the whole project uses one language (TypeScript) front to back. Express is a small web framework that maps URLs to handler functions (the REST API). '
  'Node handles many simultaneous connections efficiently, which suits a platform receiving telemetry from many devices and streaming answers to browsers. Helmet adds security headers, zod validates every input.')
h2('4.3 Why MongoDB, and why also PostgreSQL?')
bl([
    ('MongoDB for operational data.', 'Telemetry and devices are naturally documents with varying fields (different machine types report different signals). A document database stores them flexibly and scales with time-series style writes. Your faculty also asked for MongoDB, so all operational data was migrated to it.'),
    ('PostgreSQL + pgvector for RAG only.', 'The knowledge base needs vector similarity search. pgvector adds a vector column type, an HNSW index and cosine distance to PostgreSQL, plus full-text search in the same database. Keeping it for RAG gave a proven vector engine without writing one.'),
    ('Why two databases?', 'Each is used for what it is best at. The boundary is clean: MongoDB = platform data, PostgreSQL = documents and chats.'),
])
h2('4.4 Why Redis')
p('Redis is an in-memory store. It holds the rolling window of the last 20 readings per signal (needed on every new reading, so it must be very fast), rate-limit counters and the per-chat lock. '
  'Data in Redis is disposable: it can be rebuilt from MongoDB.')
h2('4.5 Why Docker, and how MongoDB was brought into it')
p('Docker packages each program with everything it needs into a container, so "it works on my machine" becomes "it works the same everywhere". One command - docker compose up -d --build - starts all five services in the right order.')
bl([
    ('Image.', 'For MongoDB no code was written: the service uses the official image mongo:7 in docker-compose.yml.'),
    ('Credentials.', 'MONGO_INITDB_ROOT_USERNAME and MONGO_INITDB_ROOT_PASSWORD create the root user on first start (defaults iot / iot for the demo, overridable through .env).'),
    ('Persistence.', 'A named volume mongodata is mounted at /data/db, so data survives restarts and container re-creation.'),
    ('Health check.', 'A mongosh ping every 5 seconds; the backend declares depends_on with condition service_healthy, so it starts only when MongoDB is really ready.'),
    ('Networking.', 'Inside the compose network the backend reaches the database by the service name: MONGO_URL = mongodb://iot:...@mongo:27017/iotmonitor?authSource=admin. The port is published only on 127.0.0.1, so nobody else on the network can reach it.'),
    ('Resource limits.', 'mem_limit on each service keeps a small laptop usable.'),
    ('Code side.', 'backend/src/mongo.ts opens the connection and creates the indexes at start-up; services use the MongoDB driver collections instead of the earlier SQL queries. Tests use throw-away databases so real data is never touched.'),
])
h2('4.6 Why a local model (Ollama) and which models')
bl([
    ('Ollama', 'is a program that runs open-source language models on your own machine. Privacy (documents never leave the computer), no cost per question, works offline, and the course project is not dependent on an external API key.'),
    ('Chat model.', 'The default is qwen2.5:7b. On this laptop the graphics card has 6 GB of memory and the 7B model did not fit, so the machine runs phi4-mini, a smaller model that fits. The model is a setting (LLM_MODEL), not hard-coded, and the code sits behind a provider interface, so it can be changed without touching the rest.'),
    ('Embedding model.', 'qwen3-embedding:0.6b produces a 1024-number vector for a piece of text. Similar meanings give nearby vectors. It is small, multilingual and fast, which is why it was chosen for retrieval.'),
    ('Trade-off to admit.', 'Small models are weaker than large cloud models and can make mistakes. The design therefore does not trust the model: it limits it to supplied sources, verifies citations, and refuses when there is no support.'),
])
h2('4.7 Security basics')
bl([
    'Passwords hashed with bcrypt; sessions use signed JWT tokens (12 hours).',
    'Roles: administrator and viewer. The first account on a fresh installation becomes administrator; later sign-ups are viewers until an admin promotes them.',
    'Devices authenticate with per-device API keys stored only as hashes.',
    'Rate limiting, Helmet headers, CORS limited to the dashboard origin, request validation with zod.',
    'Webhooks cannot target internal addresses (protection against server-side request forgery).',
    'Uploads are checked by content, not only by extension: PDF header, DOCX structure with a zip-bomb check, valid UTF-8 for text.',
    'Every sensitive action is written to an audit log.',
])

# ------------------------------------------------------------------------------------------------------------------
h1('5. The prediction model (anomaly prediction)')
bl([
    ('Question answered.', 'Will this signal cross its limit within the next 20 readings?'),
    ('Features.', 'From the last 20 readings a least-squares line gives level, slope and noise. Five features: headroom to the limit, projected movement, projected margin, volatility and trend divided by headroom.'),
    ('Algorithm.', 'Logistic regression: a weighted sum of the features passed through a sigmoid gives a probability between 0 and 1.'),
    ('Why logistic regression.', 'Fast, runs inside Node without Python, needs little data, and above all it is explainable: each feature\'s push on the result is exactly weight times value, so the "Why this risk?" panel shows exact contributions, not a guess.'),
    ('Training.', 'A global model is trained offline on synthetic drift / leak / recovery episodes. Per-machine-type models can be retrained from stored telemetry.'),
    ('Correct evaluation.', 'Time-ordered split (oldest 80 percent train, newest 20 percent test, never shuffled), with a purge gap so no test label leaks into training. Metrics: AUC, precision, recall, F1, accuracy, Brier score, average lead time.'),
    ('Validation gate.', 'A retrained model becomes active only if AUC is at least 0.70 and not more than 0.03 worse than the current model. Every run is versioned, and an administrator can roll back.'),
    ('Predictive alert rule.', 'Opens only if probability is above 50 percent AND the linear time-to-breach estimate agrees; otherwise the signal is labelled elevated. This reduces false alarms.'),
])

# ------------------------------------------------------------------------------------------------------------------
h1('6. RAG: how the AI Assistant works')
p('RAG means Retrieval-Augmented Generation: instead of letting a language model answer from memory (where it may invent things), we first RETRIEVE relevant text from our own data, '
  'then ask the model to write the answer using only that text, and show where each statement came from.')
h2('6.1 Ingestion (how a document becomes searchable)')
bl([
    '1. Upload through the Knowledge Base page (PDF, DOCX, TXT, MD, CSV); checked by content; duplicates rejected by SHA-256.',
    '2. A queue worker (FOR UPDATE SKIP LOCKED) extracts text: PDFs by page, DOCX with mammoth, CSV as "column: value" lines.',
    '3. Text is cleaned (control and direction-override characters removed).',
    '4. Chunking: about 1000 characters per chunk with 150 overlap, splitting at paragraph, sentence, then word boundaries; each chunk remembers its page and heading.',
    '5. Each chunk is turned into a 1024-dimension vector by the embedding model and stored in PostgreSQL with a full-text index column.',
    '6. The project\'s own documentation files are indexed too, so the assistant can explain the platform.',
])
h2('6.2 Retrieval (how the right text is found)')
bl([
    'Vector search: the question is embedded and the nearest chunks by cosine distance are found using the pgvector HNSW index.',
    'Keyword search: PostgreSQL full-text search catches exact terms and codes that embeddings can miss.',
    'Fusion: reciprocal rank fusion (constant 60) merges the two ranked lists - hybrid search.',
    'Relevance gate: a chunk must reach a minimum similarity (0.40) and be close to the best match, or be a real keyword match. Unrelated questions therefore retrieve nothing.',
    'Access control: the SQL filters by owner, shared or administrator BEFORE ranking, so another user\'s private document can never reach a prompt.',
    'At most 6 chunks and 6000 characters go into the prompt.',
])
h2('6.3 Live-data tools (the "database" half of the answer)')
p('Questions like "how many devices are online?" or "why is device 12 high risk?" are answered from live data. A deterministic planner picks up to four of twelve fixed tools '
  '(fleet_status, device_profile, telemetry_history, top_risk, explain_prediction, anomaly_history, maintenance_history, maintenance_attention, highest_metric, alerts_summary, breaches_by_type, model_info). '
  'Each tool is a fixed parameterised query. The model never writes SQL and never chooses a tool, so it cannot reach data the rules did not select.')
h2('6.4 Preventing hallucination (what you should emphasise)')
bl([
    'System prompt: answer ONLY from numbered sources; if they do not contain the answer say so; never invent device data, dates, thresholds or document contents.',
    'No-context shortcut: when nothing relevant is found, a fixed reply is returned and the model is not called at all.',
    'Citation verification: every [D1] / [S2] in the answer is checked against the real sources; invalid ones are removed and counted.',
    'Prompt-injection resistance: document text is untrusted data placed between markers; fake markers and fake citations inside documents are neutralised; the assistant is read-only and cannot change anything.',
    'Low temperature (0.2) and a limited context window keep answers focused.',
    'NEW - file-only questions: when you say "in the devices file that I sent" or "according to the document I uploaded", the assistant now skips live data, searches only your files, keeps only excerpts that actually share a meaningful word with your question, and if nothing matches it replies "That is not in the file(s) I can read" instead of guessing. It also ignores earlier answers in the chat for such questions so old answers cannot leak in. If the model cites nothing, the reply is replaced by the same "not in the file" message.',
])
h2('6.5 Why the assistant "hallucinated" in your screenshot')
p('You asked about a file and then "what is a cell?". The old code retrieved loosely similar excerpts from the devices file and the small model, instead of saying they did not contain the answer, '
  're-used the device list. A small model does this easily. The fix is not a better prompt alone; it is code that decides before the model is called whether there is any relevant source, and replaces uncited answers with "not in the file".')

# ------------------------------------------------------------------------------------------------------------------
h1('7. Features, screens and the demo flow')
bl([
    ('Dashboard', 'fleet health, live alerts, predicted risks.'),
    ('Devices / Device detail', 'inventory with status and firmware; live readings and forecast charts; the "Why this risk?" explanation.'),
    ('Alerts and Alert Rules', 'threshold, heartbeat and predictive alerts; limits per machine type.'),
    ('Anomaly Prediction', 'fleet risk ranking, forecast with confidence band, retrain, evaluate, versions, rollback.'),
    ('Firmware, Security Incidents, Machine Types', 'compliance and roll-outs, an incident register, per-type configuration.'),
    ('Reports', 'eight report types streamed as CSV and PDF.'),
    ('Knowledge Base and AI Assistant', 'upload documents, ask questions with citations, export a conversation.'),
    ('Users, Notifications, Audit Log, Activity', 'roles, alerts to users, who-did-what, event history.'),
])
h2('Suggested 5-minute demo')
bl([
    '1. Open the dashboard: live fleet, simulator running (say it is simulated).',
    '2. Open one device: readings, forecast chart, "Why this risk?" with exact feature contributions.',
    '3. Show an alert rule override for one machine type, then the Alerts page.',
    '4. Anomaly Prediction: risk ranking, retrain, validation gate result, version history.',
    '5. Knowledge Base: upload a short manual; ask a question about it and show the citation chip and the source excerpt.',
    '6. Ask something that is NOT in the manual and show that it says it does not have it - this is your strongest point.',
    '7. Show Users (two roles) and the Audit Log.',
    '8. Show docker compose ps (five healthy containers) and the green tests/CI.',
])

# ------------------------------------------------------------------------------------------------------------------
h1('8. Testing and quality')
bl([
    '176 backend tests (model maths, API, auth, alerts, retraining, document parsing, RAG, assistant, reports), 39 frontend tests and a 13-step end-to-end browser scenario with Playwright.',
    'Tests use throw-away databases in MongoDB and PostgreSQL and Redis database 1, never your data.',
    'A labelled deterministic test provider replaces the model in CI, so retrieval, citations and security paths are tested without a GPU.',
    'GitHub Actions CI runs typecheck, lint, tests and build for backend and frontend.',
    'TypeScript, oxlint and zod validation catch errors early.',
])

# ------------------------------------------------------------------------------------------------------------------
h1('9. Likely viva questions with answers')
QA = [
    ('What is your project in one sentence?', 'A secure, locally run platform that monitors the health and security of IoT devices, predicts imminent limit breaches with an exact explanation, and answers maintenance questions from live data and uploaded documents with verifiable sources.'),
    ('What problem does it solve?', 'Alerts fire only after a limit is crossed, one limit does not suit every machine, silent devices go unnoticed, and manuals are hard to search during a fault. The platform predicts early, uses per-type limits, detects silence by heartbeat and makes documents searchable by an assistant.'),
    ('Why did you choose this topic?', 'Unplanned downtime is expensive and the data to prevent it already exists in the readings; modern retrieval and local language models make it possible to add a safe, explainable assistant without sending data outside.'),
    ('Who are the users?', 'Two roles: administrators (manage users, rules, models, firmware) and viewers (read-only monitoring and the assistant).'),
    ('What is RAG?', 'Retrieval-Augmented Generation: retrieve relevant text from your own data first, then have the language model answer only from it and cite it. It reduces hallucination and keeps answers verifiable.'),
    ('Which model do you use and why?', 'A local model through Ollama. Default qwen2.5:7b; on this 6 GB graphics card phi4-mini is used because the 7B model does not fit. Embeddings use qwen3-embedding:0.6b (1024 dimensions). Local means privacy, no cost and offline use.'),
    ('How does the assistant avoid hallucination?', 'It is limited to numbered sources, refuses when none are relevant (without calling the model), verifies citations, treats documents as untrusted, is read-only, and for file questions answers "not in the file" if the file does not support it.'),
    ('What is an embedding?', 'A list of numbers (1024 here) that represents the meaning of a text. Similar texts get similar vectors, so searching by meaning becomes finding nearest vectors.'),
    ('What is a vector database / pgvector?', 'pgvector adds a vector type, similarity operators and an HNSW index to PostgreSQL, so we can store embeddings next to the text and search the nearest ones.'),
    ('What is chunking and why overlap?', 'Splitting documents into pieces of about 1000 characters so each can be retrieved and fits the prompt. A 150-character overlap prevents a sentence that falls on a boundary from being lost.'),
    ('What is hybrid search?', 'Vector search (meaning) plus keyword search (exact terms), merged with reciprocal rank fusion. It finds both paraphrases and exact codes or names.'),
    ('Why does the router not let the model choose tools?', 'Security and predictability: a model that picks tools or writes queries can be tricked by text in a document. Fixed rules mean it can only reach data the rules selected.'),
    ('What is prompt injection and how do you handle it?', 'Text in a document that tries to give the model orders. We place document text between markers as untrusted data, strip fake markers and citations, instruct the model to ignore commands inside sources, and the assistant is read-only anyway.'),
    ('Why MongoDB?', 'Telemetry and devices are flexible documents, MongoDB suits write-heavy time-series style data, and it was the faculty requirement. PostgreSQL remains only for vector search.'),
    ('Why both MongoDB and PostgreSQL?', 'Each does what it does best: MongoDB for operational documents, PostgreSQL with pgvector for similarity search and full-text search over the knowledge base.'),
    ('How did you bring MongoDB into Docker?', 'As a service in docker-compose.yml using the official mongo:7 image, with root credentials from environment variables, a named volume for persistence, a mongosh health check, ports bound to localhost, and the backend connecting by the service name with depends_on service_healthy.'),
    ('Why Docker?', 'Reproducibility and one-command start: five services with fixed versions, networking, volumes, health checks and memory limits, identical on any machine.'),
    ('What is the difference between an image and a container?', 'An image is the read-only template (code plus dependencies); a container is a running instance of it.'),
    ('What is a Docker volume?', 'Storage managed by Docker that outlives containers, so database data survives restarts.'),
    ('Why is Ollama not in Docker?', 'It needs the graphics card; on Windows the simplest reliable way is to run it on the host and let the backend call it through host.docker.internal. An optional compose service exists for machines that want it containerised.'),
    ('What is React and why?', 'A component-based UI library with efficient updates; ideal for a live dashboard where data changes constantly. Combined with TypeScript, React Router, TanStack Query and Recharts.'),
    ('Why Node and Express?', 'One language across the stack, good at many concurrent connections and streaming, small and well-known framework.'),
    ('Why TypeScript?', 'Static types catch mistakes before running the code and document the data shapes between frontend and backend.'),
    ('What does Redis do here?', 'Keeps the last 20 readings per signal for instant prediction, rate-limit counters and per-chat locks.'),
    ('How does prediction work?', 'Fit a line to the last 20 readings, compute five features relative to the alert limit, apply logistic regression to get the probability of breaching within 20 readings; open a predictive alert only above 50 percent and when the linear ETA agrees.'),
    ('Why logistic regression and not deep learning?', 'It is fast, needs little data, runs in Node, and is exactly explainable. Explainability matters more than a small accuracy gain for maintenance decisions.'),
    ('How do you know the model is good?', 'On held-out synthetic data AUC is about 0.94. We state clearly this is synthetic. For real data there is retraining with a time-ordered split, a purge gap and a validation gate, so a worse model is never activated.'),
    ('What is AUC?', 'Area under the ROC curve: the probability that the model ranks a random breaching case above a random non-breaching one. 0.5 is chance, 1.0 is perfect.'),
    ('Precision versus recall?', 'Precision: of the alerts raised, how many were right. Recall: of the real breaches, how many were caught. Raising the threshold trades recall for precision.'),
    ('Why a time-ordered split?', 'Random shuffling lets the model see the future during training. Ordering by time simulates real use.'),
    ('Does the model predict failure?', 'No. It predicts threshold breach within 20 readings. Failure prediction would need failure labels the platform does not have.'),
    ('How do devices authenticate?', 'Each device has its own API key sent with each reading; only a hash is stored.'),
    ('How is the system secured?', 'bcrypt password hashing, JWT sessions, role-based access, hashed device keys, validation, rate limiting, security headers, SSRF protection on webhooks, content-checked uploads and an audit log.'),
    ('What if the model is offline?', 'The health endpoint reports it, the UI shows a clear error, the question is kept and can be retried; retrieval can fall back to keyword search; live monitoring is unaffected.'),
    ('How is it tested?', '176 backend tests, 39 frontend tests, a 13-step Playwright scenario, with throw-away databases and a deterministic model double so CI needs no GPU.'),
    ('What are the limitations?', 'Simulated fleet, synthetic training data, small local model, no OCR, a single logistic regression cannot capture interactions between signals.'),
    ('What would you add next?', 'Real device onboarding, OCR for scanned PDFs, a larger model on better hardware, per-signal interaction features, notification channels such as e-mail and SMS, and horizontal scaling.'),
    ('How does the file-only question mode work?', 'The router detects wording such as "in the file I sent". It skips live tools, searches only documents, keeps excerpts that share a meaningful word with the question (or are very close matches), prefers the file named in the question, and answers "not in the file" when nothing qualifies or the model cites nothing.'),
    ('Why did the first version answer about devices when you asked "what is a cell"?', 'Similarity search always returns the nearest chunks, even if they are unrelated, and a small model tends to use whatever it is given. The relevance gate and the explicit "not in the file" path remove that behaviour.'),
]
for qq, aa in QA:
    qa(qq, aa)

h1('10. Glossary')
bl([
    ('API', 'a set of URLs through which programs talk to the server.'),
    ('REST', 'an API style using HTTP methods (GET, POST, PUT, DELETE) on resources.'),
    ('JWT', 'a signed token that proves who you are without storing a session on the server.'),
    ('RBAC', 'role-based access control: what you may do depends on your role.'),
    ('Heartbeat', 'a periodic signal showing a device is alive; silence beyond a limit means offline.'),
    ('Telemetry', 'the measured readings sent by devices.'),
    ('Embedding', 'a numeric representation of the meaning of a text.'),
    ('Cosine similarity', 'a measure of how close two vectors point; 1 is identical direction.'),
    ('HNSW', 'a graph index that finds nearest vectors quickly.'),
    ('Token', 'a piece of text (word fragment) the language model reads and writes.'),
    ('Context window', 'how much text the model can read at once (4096 tokens here).'),
    ('Temperature', 'randomness of the model; low means more consistent.'),
    ('SSE', 'server-sent events: the way answers are streamed word by word to the browser.'),
])

d.save(OUT)
print('saved', OUT)
