# -*- coding: utf-8 -*-
"""IoT Sentinel report: cover lines not bold, nothing bold in the viva paragraph, no page number on the cover / bonafide / viva pages,
roman numerals starting at the abstract (i), arabic numerals from chapter 1, table of contents lists itself."""
import copy
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC, DST = sys.argv[1], sys.argv[2]
d = Document(SRC)
ps = d.paragraphs

# ---- 1. cover: the three college lines are regular weight
for p in ps[:16]:
    if p.text.strip() in ('(AN AUTONOMOUS INSTITUTION)', 'AFFILIATED TO ANNA UNIVERSITY,', 'CHENNAI'):
        for r in p.runs:
            r.bold = False

# ---- 2. viva paragraph: nothing bold
viva = next(p for p in ps[:40] if 'Viva Voce Examination' in p.text)
for r in viva.runs:
    r.bold = False

# ---- 3. sections: [cover, bonafide, viva] without numbers | [abstract ... table of contents] roman from i | body arabic from 1
date_p = next(p for p in ps[:40] if p.text.strip() == 'Date:')
body = d.element.body
front = next(e.pPr.find(qn('w:sectPr')) for e in body if e.tag == qn('w:p') and e.pPr is not None and e.pPr.find(qn('w:sectPr')) is not None)
first = copy.deepcopy(front)
for tag in ('w:footerReference', 'w:headerReference', 'w:pgNumType', 'w:titlePg'):
    for x in first.findall(qn(tag)):
        first.remove(x)
date_p._p.get_or_add_pPr().append(first)
for x in front.findall(qn('w:titlePg')):
    front.remove(x)
pg = front.find(qn('w:pgNumType'))
pg.set(qn('w:start'), '1')

# ---- 4. table of contents lists itself
ps = d.paragraphs
toc_head = next(p for p in ps if p.text.strip() == 'TABLE OF CONTENTS' and p.style.name == 'Normal')
_bm = 7000
s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(_bm)); s.set(qn('w:name'), 'toc')
e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(_bm))
toc_head._p.pPr.addnext(s) if toc_head._p.pPr is not None else toc_head._p.insert(0, s)
s.addnext(e)
lof_entry = next(p for p in ps if p.text.replace('\t', '').strip().startswith('LIST OF FIGURES') and p is not toc_head and 'PAGEREF' in p._p.xml)
cp = copy.deepcopy(lof_entry._p)
for t in cp.iter(qn('w:t')):
    if t.text == 'LIST OF FIGURES':
        t.text = 'TABLE OF CONTENTS'
for it in cp.iter(qn('w:instrText')):
    it.text = it.text.replace('PAGEREF lof', 'PAGEREF toc')
lof_entry._p.addnext(cp)
d.save(DST)
print('saved', DST)
